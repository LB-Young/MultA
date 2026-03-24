"""PPT 提取器"""
from pathlib import Path

from src.domain.plugins import ExtractorConfig
from src.infrastructure.logging.setup import get_logger

logger = get_logger(__name__)


class PptxExtractorConfig(ExtractorConfig):
    """PPT 提取器配置"""
    supported_extensions: list[str] = [".pptx"]


class PptxExtractor:
    """PPT 演示文稿提取器"""
    
    name = "pptx"
    supported_extensions = [".pptx"]
    
    def __init__(self, config: PptxExtractorConfig = None):
        self.config = config or PptxExtractorConfig(name="pptx")
        self._pptx = None
    
    def _load_pptx(self):
        """延迟加载 pptx 库"""
        if self._pptx is None:
            try:
                from pptx import Presentation
                self._pptx = Presentation
            except ImportError:
                raise ImportError("请安装 python-pptx: pip install python-pptx")
    
    def extract(self, file_path: Path) -> str:
        """提取 PPT 演示文稿内容"""
        try:
            self._load_pptx()
            
            prs = self._pptx(str(file_path))
            paragraphs = []
            
            for slide_num, slide in enumerate(prs.slides, start=1):
                slide_content = [f"[幻灯片 {slide_num}]"]
                
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        slide_content.append(shape.text.strip())
                    
                    # 提取表格
                    if shape.has_table:
                        for row in shape.table.rows:
                            row_text = " | ".join(cell.text.strip() for cell in row.cells)
                            slide_content.append(row_text)
                
                if len(slide_content) > 1:  # 有实际内容
                    paragraphs.extend(slide_content)
                    paragraphs.append("")
            
            content = "\n".join(paragraphs)
            logger.debug(f"成功提取 PPT 文件: {file_path}, 幻灯片数: {len(prs.slides)}")
            return content
            
        except Exception as e:
            raise ValueError(f"提取 PPT 文件失败: {file_path} - {e}")
    
    def validate_file(self, file_path: Path) -> bool:
        """验证文件是否可以被提取"""
        return file_path.suffix.lower() in self.supported_extensions
